"""Add two result slides (real 20-client runs across seeds; the timed whole run) to the review deck.

Data comes from results/real_runs/summary.json and results/journey/*.json. Run from the repo root in the VM with
~/tools-venv/bin/python deliverables/scripts/add_new_result_slides.py
(The deck's original builder input was lost, so the deck is edited in place; this script is the record of that edit.)
"""
import copy
import io
import json
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.util import Emu

ROOT = Path(__file__).resolve().parents[2]
DECK = ROOT / "deliverables/presentation/FL-IDS_Review_Presentation.pptx"
TITLES = ("Results: Scaling Up and Repeating", "Results: One Whole Run, Timed")
PIC_BOX = (2036068, 1540688, 8119863, 3000000)  # the chart slot on the template slide

rows = json.load(open(ROOT / "results/real_runs/summary.json"))
ja = json.load(open(ROOT / "results/journey/phase_a_journey.json"))
jb = json.load(open(ROOT / "results/journey/phase_b_journey.json"))

prs = Presentation(DECK)
if any(s.shapes.title is not None and s.shapes.title.text_frame.text in TITLES for s in prs.slides):
    raise SystemExit("slides already added")
template = prs.slides[17]  # "Results: Resisting Sabotage"
box = next(sh for sh in template.shapes if sh.name.startswith("Rounded Rectangle"))
try:
    accent = "#" + str(box.fill.fore_color.rgb)
except Exception:
    accent = "#1F4E79"
BLUE, ORANGE, GREY = accent, "#E08A1E", "#555555"
plt.rcParams.update({"font.size": 15, "axes.spines.top": False, "axes.spines.right": False})


def png(fig):
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=200, bbox_inches="tight", transparent=False, facecolor="white")
    plt.close(fig)
    buf.seek(0)
    return buf


# Chart 1: attack detection and false alarms by seed, 10 vs 20 clients
seeds = ["42", "43", "44"]
get = lambda n, key: [next(r[key] for r in rows if r["clients"] == n and r["run"].split("seed ")[1][:2] == s) for s in seeds]
fig, (a1, a2) = plt.subplots(1, 2, figsize=(10.8, 4.0))
x = range(3)
for off, n, col in ((-0.2, 10, BLUE), (0.2, 20, ORANGE)):
    a1.bar([i + off for i in x], get(n, "attack_macro_recall"), 0.4, color=col, label=f"{n} clients")
    a2.bar([i + off for i in x], [100 * v for v in get(n, "benign_fpr")], 0.4, color=col, label=f"{n} clients")
for a, ttl in ((a1, "Attacks caught (recall)"), (a2, "False alarms (%)")):
    a.set_xticks(list(x))
    a.set_xticklabels([f"seed {s}" for s in seeds])
    a.set_title(ttl)
a1.set_ylim(0.95, 1.0)
a1.set_xlabel("(axis starts at 0.95)", color=GREY, fontsize=11)
a2.set_ylim(0, 5)
a1.legend(loc="lower right", frameon=False, fontsize=12)
chart1 = png(fig)

# Chart 2: where the time goes, Phase A and Phase B
sa = [(s["label"], s["duration_s"]) for s in ja["stages"]]
sb = [(s["label"], s["duration_s"]) for s in jb["stages"]]
short = {
    "Federated training": "Federated training",
    "Calibration split, bootstrap classifier, client partition": "Bootstrap classifier + split",
    "Load, clean and encode the dataset": "Load data",
    "Calibrate client anomaly thresholds": "Thresholds",
    "Evaluate on the clients' test slices": "Evaluate",
    "Save the model bundle": "Save bundle",
    "Load the model bundle": "Load bundle",
    "Start the SDN bridge and controller": "Start SDN bridge",
    "Network, replay, classify and block (root process)": "Replay, classify, block",
    "Check the switch's flow table": "Check flow table",
}
fig, (b1, b2) = plt.subplots(1, 2, figsize=(10.8, 4.0), gridspec_kw={"width_ratios": [1.25, 1]})
for ax, st, ttl in ((b1, sa, f"Phase A: training, {ja['elapsed_s']:.0f} s"), (b2, sb, f"Phase B: detect and block, {jb['elapsed_s']:.1f} s")):
    names = [short.get(k, k) for k, _ in st][::-1]
    vals = [v for _, v in st][::-1]
    ax.barh(names, vals, color=BLUE)
    ax.set_title(ttl)
    ax.set_xlabel("seconds")
    for i, v in enumerate(vals):
        ax.text(v, i, f" {v:.1f}", va="center", fontsize=12)
    ax.set_xlim(0, max(vals) * 1.25)
chart2 = png(fig)

m = lambda n, key: sum(r[key] for r in rows if r["clients"] == n) / 3
text1 = (
    "Real runs, not a simulation: 10 and 20 separate client processes, three seeds each. "
    f"Attacks caught on average: {m(10, 'attack_macro_recall'):.3f} with 10 sites and {m(20, 'attack_macro_recall'):.3f} with 20; "
    f"false alarms stay at {100 * min(r['benign_fpr'] for r in rows):.1f}–{100 * max(r['benign_fpr'] for r in rows):.1f}%. "
    "One 20-site run (seed 42) dips to 0.978 and is not yet explained."
)
text2 = (
    f"A complete run, from raw data to a block rule on the switch, took about {ja['elapsed_s'] + jb['elapsed_s']:.0f} seconds. "
    "Training the shared model is almost all of it; once a model exists, spotting and blocking an attack takes seconds."
)
notes1 = (
    "Earlier we only had a one-process simulation of 20 sites. We now ran the real thing: one server and 20 separate client programs, "
    "on a virtual machine with enough memory. We repeated it with three different random seeds for both 10 and 20 sites. "
    "Attack detection stays very high; one 20-site run at seed 42 is lower at 0.978, and we have not yet found out why. "
    "Only three seeds, so we do not claim confidence intervals."
)
notes2 = (
    "This is one complete run on 10 sites with the data features already prepared. The shared model takes a few minutes to train. "
    "Afterwards the live part, replaying the attack, classifying it and installing the drop rule, takes seconds."
)


def build(title, picture, text, notes):
    s = prs.slides.add_slide(template.slide_layout)
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    for sh in template.shapes:
        if sh.shape_type == 13 and sh.name != "Picture 1":  # the chart slot
            s.shapes.add_picture(picture, *[Emu(v) for v in PIC_BOX])
        elif sh.shape_type == 13:  # the logo
            s.shapes.add_picture(io.BytesIO(sh.image.blob), sh.left, sh.top, sh.width, sh.height)
        else:
            s.shapes._spTree.append(copy.deepcopy(sh._element))
    s.shapes.title.text_frame.paragraphs[0].runs[0].text = title
    for extra in s.shapes.title.text_frame.paragraphs[0].runs[1:]:
        extra.text = ""
    body = next(sh for sh in s.shapes if sh.name.startswith("Rounded Rectangle"))
    paras = body.text_frame.paragraphs
    head, rest = paras[0], paras[1]
    head.runs[0].text = "What it shows"
    rest.runs[0].text = text
    for extra in rest.runs[1:]:
        extra.text = ""
    s.notes_slide.notes_text_frame.text = notes
    return s


s1 = build(TITLES[0], chart1, text1, notes1)
s2 = build(TITLES[1], chart2, text2, notes2)
# order: scaling after "Resisting Sabotage" (slide 18), timing after the live demo (slide 19)
lst = prs.slides._sldIdLst
ids = list(lst)
new1, new2 = ids[-2], ids[-1]
lst.remove(new1)
lst.remove(new2)
lst.insert(18, new1)  # becomes slide 19
lst.insert(20, new2)  # after the live demo, now slide 20 -> timing is slide 21
prs.save(DECK)
print("slides:", len(prs.slides))
